"""导出当前分镜及生成时的产品/脚本快照。"""

from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter

from excel_export import BODY_FONT, HEADER_FILL, HEADER_FONT, STYLE_NAMES, _write_sheet, make_excel_bytes
from storyboard import CHARACTER_RULE, MODES, PRODUCT_RULE, fields_for, validate_shots


def make_storyboard_excel(record: dict) -> bytes:
    shots = validate_shots(record["shots"], record["mode"], record["inputs"][3])
    product, points, audience, duration = record["inputs"]
    workbook = load_workbook(BytesIO(make_excel_bytes(product, points, audience, duration, record["scripts"])))
    _write_sheet(workbook.create_sheet("制作说明"), [
        ("选定脚本", STYLE_NAMES[record["style"]]),
        ("制作模式", record["mode"]),
        ("分镜来源", record["generation_mode"]),
        ("已确认产品资料", record["reference"] or "未提供，请以产品实物核对"),
        ("人物设定", record["character"] or "未提供，优先使用不露脸手部镜头"),
        ("人物一致性基础要求", CHARACTER_RULE),
        ("产品一致性基础要求", PRODUCT_RULE),
        ("使用说明", "脚本和产品信息为生成此分镜时的快照，分镜为最新手动编辑版本。拍摄动作是策划，不代表实际测试；发布前核对事实及口播时长。"),
        ("AI制作说明", "提示词不能保证视频模型准确还原产品，请提供参考图并逐镜复核。" if record["mode"] == MODES[1] else "按实际拍摄条件调整机位和素材。"),
    ])
    sheet = workbook.create_sheet("视频分镜")
    fields = fields_for(record["mode"])
    sheet.append(["镜头", "时间段（秒）", *fields.values()])
    elapsed = 0
    for index, shot in enumerate(shots, 1):
        sheet.append([index, f"{elapsed}–{elapsed + shot['duration']}", *[shot[field] for field in fields]])
        elapsed += shot["duration"]
    widths = [8, 15, 14, 15, 20, 42, 42, 35] + ([65, 42, 42] if record["mode"] == MODES[1] else [])
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet:
        for cell in row:
            cell.font = HEADER_FONT if cell.row == 1 else BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == 1:
                cell.fill = HEADER_FILL
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F4FA")
            if isinstance(cell.value, str):
                cell.data_type = "s"  # 任何以 = 开头的编辑内容也只能作为文本。
        sheet.row_dimensions[row[0].row].height = 30 if row[0].row == 1 else 180
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
