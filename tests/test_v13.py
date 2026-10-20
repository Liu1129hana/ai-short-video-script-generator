"""V1.3 离线回归：真实界面流程 + 模拟 HTTP，不读取/调用真实凭据。"""

from contextlib import ExitStack
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import time
import unittest
from unittest.mock import patch

import httpx
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from deepseek_service import InvalidAIResponse
from excel_export import get_voiceover
from request_limiter import RateLimitError
from storyboard import (MODES, check_storyboard_claims, generate_storyboard, make_example_storyboard, validate_shots)
from storyboard_export import make_storyboard_excel

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = "【0—3 秒｜开头钩子】出门前看看。\n【3—9 秒｜核心卖点】杯子轻便。\n【9—16 秒｜使用场景】带去办公室。\n【16—20 秒｜购买引导】查看商品详情。\n【完整口播文案】出门前看看。杯子轻便。带去办公室。查看商品详情。"


def sample_record(mode=MODES[0], duration=20):
    return {"id": "test-record", "style": "pain", "mode": mode,
            "inputs": ("咖啡杯", "轻便", "通勤用户", duration),
            "scripts": {key: SCRIPT for key in ("pain", "experience", "promotion")},
            "reference": "", "character": "", "generation_mode": "AI生成模式",
            "shots": make_example_storyboard(SCRIPT, duration, mode)}


class FakeClient:
    def __init__(self, shots, finish_reason="stop", content=None):
        self.shots, self.finish_reason, self.content = shots, finish_reason, content
        self.calls = []

    def post(self, url, **kwargs):
        self.calls.append((url, kwargs))
        return httpx.Response(200, request=httpx.Request("POST", url), json={"choices": [{
            "finish_reason": self.finish_reason,
            "message": {"content": self.content if self.content is not None else json.dumps({"shots": self.shots}, ensure_ascii=False)},
        }]})


def call_service(client, mode=MODES[0], **kwargs):
    return generate_storyboard("咖啡杯", "轻便", "通勤用户", 20, SCRIPT, mode,
                               kwargs.get("reference", ""), "", "仅测试用占位符", "test-model", client=client)


class StoryboardServiceTests(unittest.TestCase):
    def test_common_model_field_variants_and_unused_metadata_are_normalized(self):
        aliases = {"duration": "镜头时长（秒）", "shot_size": "景别", "camera": "运镜",
                   "action": "画面动作", "dialogue": "台词", "materials": "素材需求"}
        original = sample_record()["shots"]
        returned = [{aliases[key]: f"{value}秒" if key == "duration" else value
                     for key, value in shot.items()} | {"镜头编号": index}
                    for index, shot in enumerate(original, 1)]
        content = json.dumps({"shots": returned, "summary": "拍摄说明"}, ensure_ascii=False)
        self.assertEqual(call_service(FakeClient(original, content=content)), original)
        del returned[0]["素材需求"]
        with self.assertRaisesRegex(InvalidAIResponse, "第 1 个镜头缺少“素材需求”"):
            call_service(FakeClient(original, content=json.dumps({"shots": returned}, ensure_ascii=False)))

    def test_both_modes_and_all_durations(self):
        for mode in MODES:
            for duration in (15, 20, 30, 60):
                shots = make_example_storyboard(SCRIPT, duration, mode)
                self.assertEqual(sum(item["duration"] for item in shots), duration)
                self.assertEqual("".join(item["dialogue"] for item in shots), get_voiceover(SCRIPT))
            client = FakeClient(sample_record(mode)["shots"])
            result = call_service(client, mode)
            self.assertEqual(result, sample_record(mode)["shots"])
            url, kwargs = client.calls[0]
            self.assertEqual(url, "https://api.deepseek.com/chat/completions")
            self.assertEqual(kwargs["json"]["response_format"], {"type": "json_object"})
            self.assertEqual(kwargs["json"]["model"], "test-model")
            self.assertIn("唯一产品事实依据", kwargs["json"]["messages"][0]["content"])

    def test_invalid_shape_duration_and_response_rejected(self):
        original = sample_record()["shots"]
        bads = [None, [], original[:2]]
        for field, value in (("duration", 0), ("duration", True), ("duration", 2.5), ("duration", 10), ("action", ""), ("camera", 12)):
            bad = deepcopy(original)
            bad[0][field] = value
            bads.append(bad)
        bad = deepcopy(original)
        bad[0]["extra"] = "ignored?"
        with self.assertRaises(InvalidAIResponse):
            validate_shots(bad, MODES[0], 20)  # 编辑/导出继续只接收标准字段。
        for shots in bads:
            with self.subTest(shots=shots), self.assertRaises(InvalidAIResponse):
                call_service(FakeClient(shots))
        for client in (FakeClient(original, finish_reason="length"), FakeClient(original, content="not-json"), FakeClient(original, content='{"wrong": []}')):
            with self.assertRaises(InvalidAIResponse):
                call_service(client)

    def test_generated_dialogue_cannot_silently_rewrite_current_script(self):
        shots = sample_record()["shots"]
        shots[0]["dialogue"] = "新的台词"
        with self.assertRaisesRegex(InvalidAIResponse, "改写或遗漏"):
            call_service(FakeClient(shots))

    def test_fabricated_claims_rejected_in_all_text_fields(self):
        for field in ("action", "materials", "prompt", "character", "product"):
            for claim in ("钓力12kg", "12公斤", "IP68防水", "航空铝", "实测", "仅需9.9元", "全金属", "十二轴承"):
                shots = sample_record(MODES[1])["shots"]
                shots[0][field] += claim
                with self.subTest(field=field, claim=claim), self.assertRaises(InvalidAIResponse):
                    call_service(FakeClient(shots), MODES[1])
        check_storyboard_claims("钓力12kg，杯子容量500ml。", "钓力12kg；容量500ml")
        with self.assertRaises(InvalidAIResponse):
            check_storyboard_claims("容量500kg", "容量500ml")

    def test_current_script_is_not_an_authority_for_new_parameters(self):
        with self.assertRaises(InvalidAIResponse):
            generate_storyboard("咖啡杯", "轻便", "通勤用户", 20, SCRIPT + "容量500ml。", MODES[0],
                                "", "", "test", "test-model", client=FakeClient([]))

    def test_export_contains_latest_edits_and_original_context(self):
        for mode in MODES:
            record = sample_record(mode)
            record["shots"][0]["action"] = "=HYPERLINK(\"https://example.com\")"
            record["shots"][1]["dialogue"] = "我手动写的台词。"
            book = load_workbook(BytesIO(make_storyboard_excel(record)))
            self.assertEqual(book.sheetnames, ["产品信息", "痛点营销型", "实战体验型", "促销转化型", "制作说明", "视频分镜"])
            self.assertEqual(book["视频分镜"]["F2"].value, record["shots"][0]["action"])
            self.assertEqual(book["视频分镜"]["F2"].data_type, "s")
            self.assertEqual(book["视频分镜"]["G3"].value, "我手动写的台词。")
            self.assertEqual(book["视频分镜"].max_column, 11 if mode == MODES[1] else 8)
            self.assertEqual(book["产品信息"]["B2"].value, "咖啡杯")
        record["shots"][0]["duration"] = 1
        with self.assertRaises(InvalidAIResponse):
            make_storyboard_excel(record)


class StoryboardUITests(unittest.TestCase):
    def setUp(self):
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.stack.enter_context(patch("access_control.get_app_password", return_value="测试密码"))
        self.stack.enter_context(patch("deepseek_service.get_api_key", return_value="仅测试占位符"))
        # 模块已加载时，独立绑定也要覆盖，确保绝不触及真实服务。
        self.stack.enter_context(patch("storyboard_ui.get_app_password", return_value="测试密码"))
        self.stack.enter_context(patch("storyboard_ui.get_api_key", return_value="仅测试占位符"))
        self.api = self.stack.enter_context(patch("storyboard_ui.generate_storyboard"))
        self.reserve = self.stack.enter_context(patch("storyboard_ui.reserve_ai_request"))
        self.app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=15).run()
        self.app.text_input(key="product_name").set_value("咖啡杯")
        self.app.text_area(key="selling_points").set_value("轻便")
        self.app.text_input(key="target_audience").set_value("通勤用户")
        self.app.selectbox(key="video_duration").set_value(20)
        self.app.button(key="generate_scripts").click().run()
        self.assertFalse(self.app.exception)

    def unlock(self):
        self.app.radio(key="generation_mode").set_value("AI生成模式").run()
        self.app.text_input(key="app_password_input").set_value("测试密码")
        self.app.button(key="unlock_ai").click().run()

    def generate(self):
        self.app.button(key="generate_storyboard").click().run()
        self.assertFalse(self.app.exception)

    def test_mode_and_script_switches_retain_edits_and_export_snapshot(self):
        self.generate()
        original_id = self.app.session_state["storyboards"]["pain:真人实拍"]["id"]
        key = f"shot_{original_id}_0_action"
        self.app.text_area(key=key).set_value("我改过的拍摄动作").run()
        self.app.radio(key="storyboard_mode").set_value(MODES[1]).run()
        self.generate()
        self.app.selectbox(key="storyboard_style").set_value("experience").run()
        self.generate()
        self.app.selectbox(key="storyboard_style").set_value("pain").run()
        self.app.radio(key="storyboard_mode").set_value(MODES[0]).run()
        self.assertEqual(self.app.text_area(key=key).value, "我改过的拍摄动作")
        snapshot = self.app.session_state["storyboard_export_state"].snapshot
        self.assertEqual(snapshot["record"]["shots"][0]["action"], "我改过的拍摄动作")
        self.assertEqual(len(self.app.session_state["storyboards"]), 3)
        self.assertEqual(self.api.call_count, 0)
        self.assertEqual(self.reserve.call_count, 0)

    def test_ai_mode_has_local_fallback_without_key_or_quota_use(self):
        self.app.radio(key="generation_mode").set_value("AI生成模式").run()
        self.assertTrue(self.app.button(key="generate_storyboard").disabled)
        self.app.button(key="example_storyboard_fallback").click().run()
        self.assertFalse(self.app.exception)
        record = self.app.session_state["storyboards"]["pain:真人实拍"]
        self.assertEqual(record["generation_mode"], "本地示例（基于当前脚本）")
        self.assertEqual(self.api.call_count, 0)
        self.assertEqual(self.reserve.call_count, 0)

    def test_changes_to_source_keep_board_and_flag_stale(self):
        self.generate()
        record = deepcopy(self.app.session_state["storyboards"]["pain:真人实拍"])
        self.app.text_area(key="script_pain").set_value(SCRIPT).run()
        self.app.text_input(key="product_name").set_value("另一件产品").run()
        self.app.text_area(key="storyboard_reference").set_value("白色外观").run()
        self.assertTrue(any("已变化" in item.value for item in self.app.warning))
        self.assertEqual(self.app.session_state["storyboards"]["pain:真人实拍"], record)
        self.app.button(key="generate_scripts").click().run()
        self.assertEqual(self.app.session_state["storyboards"]["pain:真人实拍"], record)

    def test_invalid_edits_disable_export_without_discarding(self):
        self.generate()
        record = self.app.session_state["storyboards"]["pain:真人实拍"]
        self.app.number_input(key=f"shot_{record['id']}_0_duration").set_value(1).run()
        self.assertFalse(self.app.session_state["storyboard_export_state"].snapshot["valid"])
        self.assertTrue(any("总时长" in item.value for item in self.app.warning))
        self.app.number_input(key=f"shot_{record['id']}_0_duration").set_value(3).run()
        self.assertTrue(self.app.session_state["storyboard_export_state"].snapshot["valid"])

    def test_auth_logout_and_password_rotation_fail_closed(self):
        self.app.radio(key="generation_mode").set_value("AI生成模式").run()
        self.assertTrue(self.app.button(key="generate_storyboard").disabled)
        self.unlock()
        self.assertFalse(self.app.button(key="generate_storyboard").disabled)
        self.app.button(key="lock_ai").click().run()
        self.assertTrue(self.app.button(key="generate_storyboard").disabled)
        self.unlock()
        with patch("storyboard_ui.get_app_password", return_value="更换后的密码"):
            self.generate()  # 已显示按钮在服务器再次验证登录状态。
        self.assertEqual(self.api.call_count, 0)
        self.assertEqual(self.reserve.call_count, 0)

    def test_uses_latest_selected_script_and_same_limiter(self):
        self.app.text_area(key="script_pain").set_value(SCRIPT).run()
        self.unlock()
        self.api.return_value = sample_record()["shots"]
        self.generate()
        self.assertEqual(self.api.call_args.args[4], SCRIPT)
        self.assertEqual(self.api.call_count, 1)
        self.assertEqual(self.reserve.call_count, 1)
        self.assertEqual(self.reserve.call_args.args[0], self.app.session_state["ai_session_id"])
        self.generate()
        self.assertEqual(self.api.call_count, 1)
        self.assertTrue(any("稍等几秒" in item.value for item in self.app.error))

    def test_failure_quota_and_missing_key_preserve_existing_board(self):
        self.generate()
        original = deepcopy(self.app.session_state["storyboards"]["pain:真人实拍"])
        self.unlock()
        self.api.side_effect = httpx.ReadTimeout("不得显示的凭据内容")
        self.generate()
        self.assertEqual(self.reserve.call_count, 1)
        self.assertEqual(self.app.session_state["storyboards"]["pain:真人实拍"], original)
        self.assertTrue(any("网络" in item.value for item in self.app.error))
        self.assertTrue(all("不得显示" not in item.value for item in self.app.error))
        self.app.session_state["last_storyboard_time"] = time.monotonic() - 10
        self.reserve.side_effect = RateLimitError("今日用量已满")
        self.generate()
        self.assertEqual(self.api.call_count, 1)
        self.assertEqual(self.app.session_state["storyboards"]["pain:真人实拍"], original)
        with patch("storyboard_ui.get_api_key", return_value=""):
            self.generate()
        self.assertEqual(self.reserve.call_count, 2)
        self.assertEqual(self.api.call_count, 1)


if __name__ == "__main__":
    unittest.main()
