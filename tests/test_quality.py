"""离线检查：绝不调用真实 DeepSeek API。"""

import json
import os
import time
import unittest
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import httpx
from streamlit.testing.v1 import AppTest
from openpyxl import load_workbook

from access_control import check_password, get_app_password, is_ai_authenticated, password_tag
from excel_export import make_excel_bytes
from deepseek_service import (
    InvalidAIResponse,
    check_unprovided_claims,
    describe_ai_error,
    generate_ai_scripts,
    get_api_key,
    get_model_name,
    parse_ai_response,
    voiceover_timing_note,
)
from request_limiter import DAILY_LIMIT, SESSION_DAILY_LIMIT, RateLimitError, reserve_ai_request


def sample_result():
    return {
        style: {
            "hook": "上班路上带咖啡，担心杯子会漏吗？",
            "selling_points": "这款杯子防漏、轻便，也容易清洗。",
            "usage_scene": "放进通勤包，到办公室再取出来喝。",
            "call_to_action": "到商品页核对详情和实际价格。",
            "voiceover": f"{style}：上班路上带咖啡，担心杯子会漏吗？这款杯子防漏、轻便，也容易清洗。放进通勤包，到办公室再取出来喝。到商品页核对详情和实际价格。",
        }
        for style in ("pain", "experience", "promotion")
    }


class QualityTests(unittest.TestCase):
    def test_password_only_from_secrets_and_compared_without_echo(self):
        with patch.dict(os.environ, {"APP_PASSWORD": "环境变量不能解锁"}):
            self.assertEqual(get_app_password({}), "")
        self.assertEqual(get_app_password({"APP_PASSWORD": ""}), "")
        self.assertEqual(get_app_password({"APP_PASSWORD": "测试用访问密码"}), "测试用访问密码")
        self.assertFalse(check_password("错误密码", "测试用访问密码"))
        self.assertTrue(check_password("测试用访问密码", "测试用访问密码"))
        self.assertFalse(is_ai_authenticated(None, "测试用访问密码"))
        self.assertTrue(is_ai_authenticated(password_tag("测试用访问密码"), "测试用访问密码"))
        self.assertFalse(is_ai_authenticated(password_tag("测试用访问密码"), "新密码"))

    def test_unauthenticated_visitor_cannot_call_paid_api(self):
        with patch("access_control.get_app_password", return_value="测试用访问密码"), patch(
            "deepseek_service.get_api_key", return_value="假密钥"
        ), patch("deepseek_service.generate_ai_scripts") as generate, patch(
            "request_limiter.reserve_ai_request"
        ) as reserve:
            app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=15).run()
            app.radio(key="generation_mode").set_value("AI生成模式").run()
            self.assertTrue(app.button(key="generate_scripts").disabled)
            self.assertFalse(app.exception)
            app.text_input(key="app_password_input").set_value("错误密码").run()
            app.button(key="unlock_ai").click().run()
            self.assertTrue(app.button(key="generate_scripts").disabled)
            self.assertTrue(any("访问密码不正确" in item.value for item in app.error))
            self.assertEqual(app.text_input(key="app_password_input").value, "")
            self.assertEqual(generate.call_count, 0)
            self.assertEqual(reserve.call_count, 0)

            app.text_input(key="app_password_input").set_value("测试用访问密码").run()
            app.button(key="unlock_ai").click().run()
            self.assertFalse(app.button(key="generate_scripts").disabled)
            self.assertFalse(app.exception)
            app.button(key="lock_ai").click().run()
            self.assertTrue(app.button(key="generate_scripts").disabled)
            self.assertEqual(generate.call_count, 0)
            self.assertEqual(reserve.call_count, 0)

    def test_missing_access_password_fails_closed(self):
        with patch("access_control.get_app_password", return_value=""), patch(
            "deepseek_service.get_api_key", return_value="假密钥"
        ), patch("deepseek_service.generate_ai_scripts") as generate:
            app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=15).run()
            app.radio(key="generation_mode").set_value("AI生成模式").run()
            self.assertTrue(app.button(key="generate_scripts").disabled)
            self.assertTrue(any("APP_PASSWORD" in item.value for item in app.warning))
            self.assertEqual(generate.call_count, 0)

    def test_five_wrong_passwords_lock_current_session(self):
        with patch("access_control.get_app_password", return_value="测试用访问密码"), patch(
            "deepseek_service.generate_ai_scripts"
        ) as generate:
            app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=15).run()
            app.radio(key="generation_mode").set_value("AI生成模式").run()
            for _ in range(5):
                app.text_input(key="app_password_input").set_value("错误密码").run()
                app.button(key="unlock_ai").click().run()
            self.assertTrue(any("15 分钟后" in item.value for item in app.error))
            app.text_input(key="app_password_input").set_value("测试用访问密码").run()
            app.button(key="unlock_ai").click().run()
            self.assertTrue(app.button(key="generate_scripts").disabled)
            self.assertEqual(generate.call_count, 0)

    def test_offline_generation_and_schema_rejection(self):
        class FakeResponse:
            def raise_for_status(self):
                return None

            def json(self):
                return {"choices": [{
                    "finish_reason": "stop",
                    "message": {"content": json.dumps(sample_result(), ensure_ascii=False)},
                }]}

        class FakeClient:
            def post(self, url, **kwargs):
                self_outer.assertEqual(url, "https://api.deepseek.com/chat/completions")
                self_outer.assertEqual(kwargs["json"]["response_format"], {"type": "json_object"})
                self_outer.assertEqual(kwargs["json"]["model"], "test-model")
                self_outer.assertIn("JSON", kwargs["json"]["messages"][0]["content"])
                self_outer.assertEqual(kwargs["headers"], {"Authorization": "Bearer 仅用于测试的假密钥"})
                return FakeResponse()

        self_outer = self
        client = FakeClient()
        result = generate_ai_scripts(
            "便携咖啡杯", "防漏、轻便、容易清洗", "通勤上班族", 30,
            "仅用于测试的假密钥", "test-model", client=client,
        )
        self.assertEqual(set(result), {"pain", "experience", "promotion"})

        with patch.dict(os.environ, {"DEEPSEEK_MODEL": "", "DEEPSEEK_API_KEY": "", "GEMINI_API_KEY": "旧密钥"}):
            self.assertEqual(get_model_name(), "deepseek-flash")
            self.assertEqual(get_api_key({}), "")

        extra = sample_result()
        extra["other"] = extra["pain"]
        with self.assertRaises(InvalidAIResponse):
            parse_ai_response(json.dumps(extra, ensure_ascii=False))
        missing_field = sample_result()
        del missing_field["pain"]["voiceover"]
        with self.assertRaises(InvalidAIResponse):
            parse_ai_response(json.dumps(missing_field, ensure_ascii=False))
        duplicate = sample_result()
        duplicate["promotion"]["voiceover"] = duplicate["pain"]["voiceover"]
        with self.assertRaises(InvalidAIResponse):
            parse_ai_response(json.dumps(duplicate, ensure_ascii=False))

    def test_claim_and_timing_checks(self):
        scripts = sample_result()
        scripts["pain"]["voiceover"] += "今天只要9.9元。"
        with self.assertRaises(InvalidAIResponse):
            check_unprovided_claims(scripts, "便携咖啡杯 防漏、轻便、容易清洗 通勤上班族")
        check_unprovided_claims(scripts, "便携咖啡杯 防漏、轻便、容易清洗 通勤上班族 9.9元")
        self.assertIn("短于 15 秒", voiceover_timing_note("你好", 15))
        self.assertIn("超过 15 秒", voiceover_timing_note("通勤便携咖啡杯" * 20, 15))
        self.assertIsNone(voiceover_timing_note("上班路上带咖啡更轻松" * 8, 30))

    def test_error_messages_never_echo_exception(self):
        class ApiError(Exception):
            def __init__(self, code):
                self.code = code
                super().__init__("假密钥不应出现在界面")

        for code, expected in ((400, "格式"), (401, "密钥"), (402, "余额"), (403, "权限"),
                               (404, "模型"), (429, "频繁"), (503, "暂时")):
            message = describe_ai_error(ApiError(code))
            self.assertIn(expected, message)
            self.assertNotIn("假密钥", message)
        request = httpx.Request("POST", "https://api.deepseek.com/chat/completions")
        response = httpx.Response(402, request=request, text="假密钥不应出现在界面")
        message = describe_ai_error(httpx.HTTPStatusError("失败", request=request, response=response))
        self.assertIn("余额", message)
        self.assertNotIn("假密钥", message)
        self.assertIn("网络", describe_ai_error(httpx.ConnectError("连接失败")))

    def test_page_retains_edits_and_blocks_repeat_click(self):
        with patch("access_control.get_app_password", return_value="测试用访问密码"), patch(
            "deepseek_service.get_api_key", return_value="测试密钥"
        ), patch(
            "deepseek_service.generate_ai_scripts", return_value=sample_result()
        ) as generate, patch("request_limiter.reserve_ai_request") as reserve:
            app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=15).run()
            app.text_input(key="product_name").set_value("便携咖啡杯").run()
            app.text_area(key="selling_points").set_value("防漏、轻便、容易清洗").run()
            app.text_input(key="target_audience").set_value("通勤上班族").run()
            app.radio(key="generation_mode").set_value("AI生成模式").run()
            app.text_input(key="app_password_input").set_value("测试用访问密码").run()
            app.button(key="unlock_ai").click().run()
            app.button(key="generate_scripts").click().run()
            self.assertEqual(generate.call_count, 1)
            self.assertEqual(reserve.call_count, 1)
            self.assertIn("【完整口播文案】", app.text_area(key="script_pain").value)
            app.button(key="generate_scripts").click().run()
            self.assertEqual(generate.call_count, 1)
            self.assertTrue(any("避免重复调用" in item.value for item in app.warning))

        original = app.text_area(key="script_pain").value
        with patch("access_control.get_app_password", return_value="测试用访问密码"):
            app.text_area(key="script_pain").set_value(original.replace("上班路上带咖啡", "用户改过的开头", 1)).run()
        edited_scripts = {
            style: app.text_area(key=f"script_{style}").value
            for style in ("pain", "experience", "promotion")
        }
        workbook = load_workbook(BytesIO(make_excel_bytes(
            "便携咖啡杯", "防漏、轻便、容易清洗", "通勤上班族", 30,
            edited_scripts, date(2026, 10, 19),
        )))
        self.assertEqual(workbook.sheetnames, ["产品信息", "痛点营销型", "实战体验型", "促销转化型"])
        self.assertEqual(workbook["痛点营销型"]["B4"].value, "用户改过的开头，担心杯子会漏吗？")
        signature = app.session_state["last_ai_request"][0]
        app.session_state["last_ai_request"] = (signature, time.monotonic() - 20)

        class Denied(Exception):
            code = 403

        with patch("access_control.get_app_password", return_value="测试用访问密码"), patch(
            "deepseek_service.get_api_key", return_value="测试密钥"
        ), patch(
            "deepseek_service.generate_ai_scripts", side_effect=Denied("密钥内容不应显示")
        ), patch("request_limiter.reserve_ai_request"):
            app.button(key="generate_scripts").click().run()
        self.assertFalse(app.exception)
        self.assertIn("用户改过的开头", app.text_area(key="script_pain").value)
        self.assertTrue(any("没有访问权限" in item.value for item in app.error))
        self.assertTrue(all("密钥内容不应显示" not in item.value for item in app.error))

    def test_shared_daily_limit(self):
        with TemporaryDirectory() as directory:
            db_path = Path(directory) / "usage.sqlite3"
            today = datetime(2026, 10, 19, 12, 0)
            for index in range(DAILY_LIMIT):
                reserve_ai_request(f"session-{index // SESSION_DAILY_LIMIT}", db_path=db_path, now=today)
            with self.assertRaisesRegex(RateLimitError, "全站"):
                reserve_ai_request("another", db_path=db_path, now=today)
            reserve_ai_request("another", db_path=db_path, now=datetime(2026, 10, 20, 12, 0))

    def test_session_limit(self):
        with TemporaryDirectory() as directory:
            db_path = Path(directory) / "usage.sqlite3"
            for _ in range(SESSION_DAILY_LIMIT):
                reserve_ai_request("same", db_path=db_path, now=datetime(2026, 10, 19))
            with self.assertRaisesRegex(RateLimitError, "当前页面"):
                reserve_ai_request("same", db_path=db_path, now=datetime(2026, 10, 19))


if __name__ == "__main__":
    unittest.main()
