"""示例模式验收：只用本地模板，不调用 DeepSeek API。"""

import unittest
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from excel_export import make_excel_bytes


ROOT = Path(__file__).resolve().parents[1]
STYLES = {"pain": "痛点营销型", "experience": "实战体验型", "promotion": "促销转化型"}
REQUIRED_FIELDS = {
    "脚本类型", "视频时长", "前三秒钩子", "产品核心卖点",
    "实战使用场景", "结尾购买引导", "完整口播文案",
}


class AcceptanceTests(unittest.TestCase):
    def test_all_durations_three_styles_and_excel(self):
        app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=15).run()
        self.assertFalse(app.exception)
        app.text_input(key="product_name").set_value("便携咖啡杯").run()
        app.text_area(key="selling_points").set_value("防漏、轻便、容易清洗").run()
        app.text_input(key="target_audience").set_value("通勤上班族").run()

        for duration in (15, 20, 30, 60):
            with self.subTest(duration=duration):
                app.selectbox(key="video_duration").set_value(duration).run()
                app.button[0].click().run()
                self.assertFalse(app.exception)
                scripts = {
                    key: app.text_area(key=f"script_{key}").value for key in STYLES
                }
                self.assertEqual(len(set(scripts.values())), 3)
                for script in scripts.values():
                    self.assertIn("【0—3 秒｜开头钩子】", script)
                    self.assertIn(f"—{duration} 秒｜购买引导】", script)
                    for label in ("核心卖点", "使用场景"):
                        self.assertIn(f"｜{label}】", script)

                book = load_workbook(BytesIO(make_excel_bytes(
                    "便携咖啡杯", "防漏、轻便、容易清洗", "通勤上班族",
                    duration, scripts, date(2026, 10, 19),
                )))
                self.assertEqual(book.sheetnames, ["产品信息", *STYLES.values()])
                info = {row[0].value: row[1].value for row in book["产品信息"].iter_rows(min_row=2, max_col=2)}
                self.assertEqual(info["产品名称"], "便携咖啡杯")
                self.assertEqual(info["视频时长"], duration)
                for key, name in STYLES.items():
                    rows = {row[0].value: row[1].value for row in book[name].iter_rows(min_row=2, max_col=2)}
                    self.assertEqual(set(rows), REQUIRED_FIELDS)
                    self.assertTrue(all(rows.values()))
                    self.assertEqual(rows["脚本类型"], name)
                    self.assertEqual(rows["视频时长"], duration)
                    self.assertIn(rows["前三秒钩子"], rows["完整口播文案"])
                    self.assertIn(rows["结尾购买引导"], rows["完整口播文案"])
                    self.assertTrue(book[name]["A1"].font.bold)
                    self.assertTrue(book[name]["B8"].alignment.wrap_text)


if __name__ == "__main__":
    unittest.main()
