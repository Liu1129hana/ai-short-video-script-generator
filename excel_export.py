"""将当前页面的产品信息和示例脚本导出为 Excel。"""

import math
import re
import unicodedata
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


STYLE_NAMES = {
    "pain": "痛点营销型",
    "experience": "实战体验型",
    "promotion": "促销转化型",
}
SECTION_LABELS = ("开头钩子", "核心卖点", "使用场景", "购买引导")
SECTION_PATTERN = re.compile(
    r"【[^】｜]*｜\s*(开头钩子|核心卖点|使用场景|购买引导)\s*】"
)
VOICEOVER_PATTERN = re.compile(r"【\s*完整口播文案\s*】")
HEADER_FILL = PatternFill("solid", fgColor="17315B")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=11, color="233142")


def parse_sections(script: str) -> dict[str, str]:
    """按页面中的四段标签拆分脚本，始终使用编辑后的原文。"""
    sections = {label: "" for label in SECTION_LABELS}
    section_text = VOICEOVER_PATTERN.split(script, maxsplit=1)[0]
    matches = list(SECTION_PATTERN.finditer(section_text))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(section_text)
        content = section_text[match.end() : end].strip()
        label = match.group(1)
        sections[label] = "\n".join(part for part in (sections[label], content) if part)
    return sections


def missing_sections(script: str) -> list[str]:
    sections = parse_sections(script)
    missing = [label for label in SECTION_LABELS if not sections[label]]
    voiceover_match = VOICEOVER_PATTERN.search(script)
    if voiceover_match and not script[voiceover_match.end() :].strip():
        missing.append("完整口播文案")
    return missing


def get_voiceover(script: str) -> str:
    """读取用户编辑的完整口播；旧示例没有单独口播时由四段组成。"""
    voiceover_match = VOICEOVER_PATTERN.search(script)
    if voiceover_match:
        return script[voiceover_match.end() :].strip()
    return SECTION_PATTERN.sub("", script).strip()


def export_filename(product: str, export_date: date | None = None) -> str:
    """保留可读的产品名称，同时去掉文件名不能使用的字符。"""
    safe_product = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', "_", product.strip())
    safe_product = safe_product.strip(" ._")[:40] or "未命名产品"
    today = export_date or date.today()
    return f"{safe_product}_营销脚本_{today:%Y%m%d}.xlsx"


def _display_width(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(char) in "FW" else 1 for char in text)


def _write_sheet(worksheet, rows: list[tuple[str, str | int | date]]) -> None:
    worksheet.append(("字段", "内容"))
    for label, value in rows:
        worksheet.append((label, value))

    worksheet.column_dimensions["A"].width = 23
    worksheet.column_dimensions["B"].width = 74
    worksheet.row_dimensions[1].height = 30
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1
    worksheet.print_area = f"A1:B{worksheet.max_row}"

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for row_number in range(2, worksheet.max_row + 1):
        label_cell = worksheet.cell(row_number, 1)
        value_cell = worksheet.cell(row_number, 2)
        for cell in (label_cell, value_cell):
            cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
        label_cell.font = Font(name="Arial", size=11, bold=True, color="233142")
        if label_cell.value == "视频时长":
            value_cell.number_format = '0"秒"'
        elif isinstance(value_cell.value, date):
            value_cell.number_format = "yyyy-mm-dd"
        value = str(value_cell.value or "")
        line_count = sum(max(1, math.ceil(_display_width(part) / 70)) for part in value.splitlines())
        worksheet.row_dimensions[row_number].height = min(390, max(28, line_count * 18 + 10))


def make_excel_bytes(
    product: str,
    selling_points: str,
    audience: str,
    duration: int,
    scripts: dict[str, str],
    export_date: date | None = None,
) -> bytes:
    """返回 .xlsx 文件内容；脚本参数必须来自当前编辑框。"""
    for style_key in STYLE_NAMES:
        missing = missing_sections(scripts[style_key])
        if missing:
            raise ValueError(f"{STYLE_NAMES[style_key]}缺少：{'、'.join(missing)}")

    workbook = Workbook()
    info_sheet = workbook.active
    info_sheet.title = "产品信息"
    _write_sheet(
        info_sheet,
        [
            ("产品名称", product),
            ("核心卖点", selling_points),
            ("目标用户", audience),
            ("视频时长", duration),
            ("生成日期", export_date or date.today()),
        ],
    )

    for style_key, style_name in STYLE_NAMES.items():
        script = scripts[style_key]
        sections = parse_sections(script)
        voiceover = get_voiceover(script)
        worksheet = workbook.create_sheet(style_name)
        _write_sheet(
            worksheet,
            [
                ("脚本类型", style_name),
                ("视频时长", duration),
                ("前三秒钩子", sections["开头钩子"]),
                ("产品核心卖点", sections["核心卖点"]),
                ("实战使用场景", sections["使用场景"]),
                ("结尾购买引导", sections["购买引导"]),
                ("完整口播文案", voiceover),
            ],
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
